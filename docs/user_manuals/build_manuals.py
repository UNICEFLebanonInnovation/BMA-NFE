#!/usr/bin/env python3
"""Build the NFE and ALP end-user manuals (English and Arabic) as Markdown and Word files.

Usage::

    python docs/user_manuals/build_manuals.py            # Markdown + DOCX
    python docs/user_manuals/build_manuals.py --pdf      # also render PDF with LibreOffice
    python docs/user_manuals/build_manuals.py --font "Noto Naskh Arabic" --latin-font Inter

The manual text lives in ``content_nfe.py`` and ``content_alp.py``; screenshots live in
``screenshots/<lang>/``. Every section is a list of blocks:

    ('p', text)                       paragraph
    ('note', text)                    highlighted tip / warning
    ('bullets', [text, ...])          bulleted list
    ('steps', [(text, image, caption), ...])  numbered steps, image/caption optional
    ('img', image, caption)           stand-alone screenshot
    ('table', [headers], [[cells], ...])
"""
import argparse
import datetime
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import content_alp  # noqa: E402
import content_nfe  # noqa: E402

PRIMARY = RGBColor(0x00, 0x33, 0x66)
ACCENT = RGBColor(0x00, 0xAD, 0xEF)
MUTED = RGBColor(0x55, 0x5E, 0x66)

STRINGS = {
    'en': {'figure': 'Figure', 'contents': 'Contents', 'step': 'Step', 'note': 'Note',
           'page': 'Page', 'generated': 'Last updated', 'version': 'Version'},
    'ar': {'figure': 'الشكل', 'contents': 'المحتويات', 'step': 'الخطوة', 'note': 'ملاحظة',
           'page': 'صفحة', 'generated': 'آخر تحديث', 'version': 'الإصدار'},
}


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------
def build_markdown(manual, lang, out_path):
    S = STRINGS[lang]
    rel_img = '../screenshots/%s/' % lang
    fig = [0]
    lines = []
    if lang == 'ar':
        lines.append('<div dir="rtl">\n')
    lines.append('# %s\n' % manual['title'])
    lines.append('**%s**\n' % manual['subtitle'])
    lines.append('_%s: %s_\n' % (S['generated'], datetime.date.today().strftime('%d %B %Y')))
    for p in manual['intro']:
        lines.append(p + '\n')
    lines.append('## %s\n' % S['contents'])
    for i, sec in enumerate(manual['sections'], 1):
        lines.append('%d. %s' % (i, sec['title']))
    lines.append('')

    def image(img, caption):
        fig[0] += 1
        lines.append('![%s %d](%s%s)\n' % (S['figure'], fig[0], rel_img, img))
        lines.append('*%s %d – %s*\n' % (S['figure'], fig[0], caption))

    for i, sec in enumerate(manual['sections'], 1):
        lines.append('## %d. %s\n' % (i, sec['title']))
        for block in sec['body']:
            kind = block[0]
            if kind == 'p':
                lines.append(block[1] + '\n')
            elif kind == 'note':
                lines.append('> **%s:** %s\n' % (S['note'], block[1]))
            elif kind == 'bullets':
                for b in block[1]:
                    lines.append('- ' + b)
                lines.append('')
            elif kind == 'steps':
                for n, step in enumerate(block[1], 1):
                    text, img, caption = (list(step) + [None, None])[:3]
                    lines.append('%d. %s' % (n, text))
                    if img:
                        lines.append('')
                        image(img, caption or text)
                lines.append('')
            elif kind == 'img':
                image(block[1], block[2])
            elif kind == 'table':
                headers, rows = block[1], block[2]
                lines.append('| ' + ' | '.join(headers) + ' |')
                lines.append('|' + '---|' * len(headers))
                for r in rows:
                    lines.append('| ' + ' | '.join(r) + ' |')
                lines.append('')
    if lang == 'ar':
        lines.append('</div>')
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, 'w', encoding='utf-8') as fh:
        fh.write('\n'.join(lines))


# ---------------------------------------------------------------------------
# DOCX helpers
# ---------------------------------------------------------------------------
class DocBuilder:
    def __init__(self, lang, font, latin_font):
        self.lang = lang
        self.rtl = lang == 'ar'
        self.font = font
        self.latin_font = latin_font
        self.S = STRINGS[lang]
        self.fig = 0
        self.doc = Document()
        self._setup()

    # -- document defaults --------------------------------------------------
    def _setup(self):
        doc = self.doc
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ('left_margin', 'right_margin'):
            setattr(sec, side, Cm(2.0))
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
        styles = doc.styles
        normal = styles['Normal']
        normal.font.name = self.latin_font
        normal.font.size = Pt(11)
        self._set_fonts(normal.element.get_or_add_rPr())
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        if self.rtl:
            ppr = normal.element.get_or_add_pPr()
            bidi = OxmlElement('w:bidi')
            ppr.append(bidi)
        for lvl, size in ((1, 20), (2, 15), (3, 12.5)):
            st = styles['Heading %d' % lvl]
            st.font.name = self.latin_font
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = PRIMARY
            self._set_fonts(st.element.get_or_add_rPr())
            st.paragraph_format.space_before = Pt(18 if lvl == 1 else 12)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.keep_with_next = True
        # footer with page number
        footer = sec.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run(self.S['page'] + ' ')
        self._style_run(run, size=9, color=MUTED)
        self._add_field(footer, 'PAGE')

    def _set_fonts(self, rpr):
        rfonts = rpr.find(qn('w:rFonts'))
        if rfonts is None:
            rfonts = OxmlElement('w:rFonts')
            rpr.insert(0, rfonts)
        rfonts.set(qn('w:ascii'), self.latin_font)
        rfonts.set(qn('w:hAnsi'), self.latin_font)
        rfonts.set(qn('w:eastAsia'), self.latin_font)
        rfonts.set(qn('w:cs'), self.font)

    def _add_field(self, paragraph, instr):
        run = paragraph.add_run()
        self._style_run(run, size=9, color=MUTED)
        for tag, text in (('begin', None), ('instr', instr), ('separate', None), ('text', '1'), ('end', None)):
            if tag == 'instr':
                el = OxmlElement('w:instrText')
                el.set(qn('xml:space'), 'preserve')
                el.text = ' %s ' % text
            elif tag == 'text':
                el = OxmlElement('w:t')
                el.text = text
            else:
                el = OxmlElement('w:fldChar')
                el.set(qn('w:fldCharType'), tag)
            run._r.append(el)

    # -- low level formatting --------------------------------------------------
    def _style_run(self, run, size=None, bold=None, italic=None, color=None, latin=False):
        run.font.name = self.latin_font
        self._set_fonts(run._r.get_or_add_rPr())
        if latin:
            # Latin words and digits inside Arabic paragraphs: no complex-script flag
            if size:
                run.font.size = Pt(size)
            if bold is not None:
                run.font.bold = bold
            if italic is not None:
                run.font.italic = italic
            if color is not None:
                run.font.color.rgb = color
            return
        if size:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if italic is not None:
            run.font.italic = italic
        if color is not None:
            run.font.color.rgb = color
        if self.rtl:
            rpr = run._r.get_or_add_rPr()
            if rpr.find(qn('w:rtl')) is None:
                rpr.append(OxmlElement('w:rtl'))
            # complex-script size/bold so Word renders Arabic consistently
            if size:
                szcs = OxmlElement('w:szCs')
                szcs.set(qn('w:val'), str(int(size * 2)))
                rpr.append(szcs)
            if bold:
                rpr.append(OxmlElement('w:bCs'))

    def _para(self, text='', style=None, align=None, size=None, bold=None, italic=None, color=None,
              space_after=None, keep_with_next=False):
        p = self.doc.add_paragraph(style=style)
        if self.rtl:
            ppr = p._p.get_or_add_pPr()
            ppr.append(OxmlElement('w:bidi'))
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align is None else align
        elif align is not None:
            p.alignment = align
        if text:
            self._rich(p, text, size=size, bold=bold, italic=italic, color=color)
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        if keep_with_next:
            p.paragraph_format.keep_with_next = True
        return p

    LATIN = re.compile(r"[A-Za-z0-9][A-Za-z0-9 .,:/@()\-_+#!$%&*]*[A-Za-z0-9)%]|[A-Za-z0-9]")

    def _rich(self, p, text, **kw):
        """Support **bold** spans inside text; Latin/digit segments keep the Latin font in Arabic text."""
        parts = text.split('**')
        for i, part in enumerate(parts):
            if not part:
                continue
            kw2 = dict(kw)
            if i % 2 == 1:
                kw2['bold'] = True
            if not self.rtl:
                self._style_run(p.add_run(part), **kw2)
                continue
            pos = 0
            for m in self.LATIN.finditer(part):
                if m.start() > pos:
                    self._style_run(p.add_run(part[pos:m.start()]), **kw2)
                self._style_run(p.add_run(m.group(0)), latin=True, **kw2)
                pos = m.end()
            if pos < len(part):
                self._style_run(p.add_run(part[pos:]), **kw2)

    # -- blocks ---------------------------------------------------------------
    def heading(self, text, level):
        p = self.doc.add_heading(level=level)
        if self.rtl:
            ppr = p._p.get_or_add_pPr()
            ppr.append(OxmlElement('w:bidi'))
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(text)
        self._style_run(run, bold=True, color=PRIMARY, size={1: 20, 2: 15, 3: 12.5}[level])
        return p

    def paragraph(self, text):
        return self._para(text)

    def note(self, text):
        p = self._para()
        self._shade(p, 'EAF6FD')
        self._rich(p, '**%s:** ' % self.S['note'] + text, size=10.5)
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.right_indent = Cm(0.3)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)

    def _shade(self, p, hex_color):
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        ppr.append(shd)
        pbdr = OxmlElement('w:pBdr')
        side = OxmlElement('w:right' if self.rtl else 'w:left')
        side.set(qn('w:val'), 'single')
        side.set(qn('w:sz'), '18')
        side.set(qn('w:space'), '6')
        side.set(qn('w:color'), '00ADEF')
        pbdr.append(side)
        ppr.append(pbdr)

    def bullets(self, items):
        for item in items:
            p = self._para(style='List Bullet')
            self._rich(p, item)
            p.paragraph_format.space_after = Pt(3)

    def steps(self, items):
        for n, step in enumerate(items, 1):
            text, img, caption = (list(step) + [None, None])[:3]
            p = self._para(keep_with_next=bool(img))
            p.paragraph_format.left_indent = Cm(0.9) if not self.rtl else None
            p.paragraph_format.right_indent = Cm(0.9) if self.rtl else None
            p.paragraph_format.first_line_indent = Cm(-0.9)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run('%d.\t' % n)
            self._style_run(run, bold=True, color=ACCENT)
            self._rich(p, text)
            if img:
                self.image(img, caption or text)

    def image(self, img, caption):
        path = os.path.join(HERE, 'screenshots', self.lang, img)
        if not os.path.exists(path):
            print('  !! missing screenshot', path)
            return
        self.fig += 1
        p = self._para(align=WD_ALIGN_PARAGRAPH.CENTER, keep_with_next=True, space_after=2)
        with Image.open(path) as im:
            w, h = im.size
        width = Cm(16.0)
        if h / w * 16.0 > 20.0:  # very tall full-page capture: limit the height instead
            width = Cm(20.0 * w / h)
        p.add_run().add_picture(path, width=width)
        self._border(p)
        c = self._para(align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
        self._rich(c, '%s %d – %s' % (self.S['figure'], self.fig, caption), size=9.5, italic=True, color=MUTED)

    def _border(self, p):
        ppr = p._p.get_or_add_pPr()
        pbdr = OxmlElement('w:pBdr')
        for side in ('top', 'left', 'bottom', 'right'):
            el = OxmlElement('w:%s' % side)
            el.set(qn('w:val'), 'single')
            el.set(qn('w:sz'), '4')
            el.set(qn('w:space'), '4')
            el.set(qn('w:color'), 'D0D7DE')
            pbdr.append(el)
        ppr.append(pbdr)

    def table(self, headers, rows):
        table = self.doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        if self.rtl:
            tblpr = table._tbl.tblPr
            tblpr.append(OxmlElement('w:bidiVisual'))
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ''
            p = cell.paragraphs[0]
            if self.rtl:
                p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            self._rich(p, h, bold=True, size=10, color=RGBColor(0xFF, 0xFF, 0xFF))
            self._cell_shade(cell, '003366')
        for r in rows:
            cells = table.add_row().cells
            for i, val in enumerate(r):
                cells[i].text = ''
                p = cells[i].paragraphs[0]
                if self.rtl:
                    p._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                self._rich(p, val, size=10)
                p.paragraph_format.space_after = Pt(2)
        self._para(space_after=6)

    def _cell_shade(self, cell, hex_color):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcpr.append(shd)

    def page_break(self):
        self.doc.add_page_break()

    # -- whole manual --------------------------------------------------------
    def build(self, manual, out_path):
        S = self.S
        # Title page
        for _ in range(6):
            self._para()
        self._para(manual['title'], align=WD_ALIGN_PARAGRAPH.CENTER, size=30, bold=True, color=PRIMARY, space_after=10)
        self._para(manual['subtitle'], align=WD_ALIGN_PARAGRAPH.CENTER, size=15, color=ACCENT, space_after=28)
        self._para(manual['audience'], align=WD_ALIGN_PARAGRAPH.CENTER, size=12, color=MUTED, space_after=6)
        self._para('%s: %s' % (S['generated'], datetime.date.today().strftime('%d %B %Y')),
                   align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=MUTED)
        self._para('%s 1.0' % S['version'], align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=MUTED)
        self.page_break()
        # Contents (static list; headings are real Word headings so a live TOC can be inserted later)
        self.heading(S['contents'], 1)
        for i, sec in enumerate(manual['sections'], 1):
            p = self._para(space_after=2)
            self._rich(p, '%d.  %s' % (i, sec['title']), size=11)
        self.page_break()
        for ptext in manual['intro']:
            self.paragraph(ptext)
        for i, sec in enumerate(manual['sections'], 1):
            self.page_break()
            self.heading('%d. %s' % (i, sec['title']), 1)
            for block in sec['body']:
                kind = block[0]
                if kind == 'p':
                    self.paragraph(block[1])
                elif kind == 'h3':
                    self.heading(block[1], 2)
                elif kind == 'note':
                    self.note(block[1])
                elif kind == 'bullets':
                    self.bullets(block[1])
                elif kind == 'steps':
                    self.steps(block[1])
                elif kind == 'img':
                    self.image(block[1], block[2])
                elif kind == 'table':
                    self.table(block[1], block[2])
        os.makedirs(os.path.dirname(out_path), exist_ok=True)
        self.doc.save(out_path)


def to_pdf(docx_path):
    out_dir = os.path.dirname(docx_path)
    cmd = ['soffice', '--headless', '--convert-to', 'pdf', '--outdir', out_dir, docx_path]
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=600)


MANUALS = [
    ('NFE', content_nfe.MANUALS),
    ('ALP', content_alp.MANUALS),
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--pdf', action='store_true', help='also export PDF files with LibreOffice')
    ap.add_argument('--font', default='Arial', help='complex-script (Arabic) font name used in the DOCX')
    ap.add_argument('--latin-font', default='Calibri', help='Latin font name used in the DOCX')
    ap.add_argument('--out', default=HERE, help='output directory for DOCX/PDF')
    args = ap.parse_args()
    for module, manuals in MANUALS:
        for lang in ('en', 'ar'):
            manual = manuals[lang]
            md_path = os.path.join(HERE, lang, '%s_User_Manual.md' % module)
            build_markdown(manual, lang, md_path)
            docx_path = os.path.join(args.out, '%s_User_Manual_%s.docx' % (module, lang.upper()))
            DocBuilder(lang, args.font, args.latin_font).build(manual, docx_path)
            print('built', md_path, docx_path)
            if args.pdf:
                to_pdf(docx_path)
                print('  pdf ok')


if __name__ == '__main__':
    main()
