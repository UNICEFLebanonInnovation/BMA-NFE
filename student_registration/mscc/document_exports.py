"""Generate MSCC documents from the same registered child and photograph.

Each workbook row represents one registration. Child and registration IDs
remain in every export so a child registered in several rounds can still be
identified without losing the programme or round of a particular record.
"""

import csv
import datetime
import io
import json
import logging
import re
from collections import Counter, OrderedDict
from dataclasses import dataclass

import xlsxwriter
from django.db import models
from PIL import Image, ImageOps

from student_registration.child.models import Child
from student_registration.mscc.models import ExceptionalExam, Registration


logger = logging.getLogger(__name__)

PHOTO_INCLUDED = 'Photograph included'
PHOTO_MISSING = 'No photograph uploaded'
PHOTO_UNAVAILABLE = 'Photograph unavailable'

# XML 1.0 cannot represent these characters in XLSX or DOCX documents.
_INVALID_XML_CHARACTERS = re.compile(
    '[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]'
)


@dataclass(frozen=True)
class ExportColumn:
    key: str
    label: str
    aliases: tuple = ()


@dataclass(frozen=True)
class RegistrationSnapshot:
    values: OrderedDict
    photo_data: bytes | None
    photo_status: str


_REQUIRED_COLUMNS = (
    ExportColumn('registration_id', 'Registration ID', ('id',)),
    ExportColumn('child_id', 'Child ID'),
    ExportColumn('photograph', 'Photograph', ('photo',)),
    ExportColumn('photograph_status', 'Photograph status', ('photo_status',)),
)


def _text(value):
    """Return literal, XML-safe text without inventing values for empty data."""
    if value is None:
        return ''
    if isinstance(value, bool):
        value = 'Yes' if value else 'No'
    elif isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        value = value.isoformat()
    elif isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    return _INVALID_XML_CHARACTERS.sub('', str(value))


def _model_columns(model, prefix, excluded):
    fields = [field for field in model._meta.concrete_fields if field.name not in excluded]
    labels = [_text(field.verbose_name) for field in fields]
    duplicates = Counter(labels)
    columns = []
    for field, label in zip(fields, labels):
        # Several legacy fields have the same verbose name. Keep each value
        # identifiable while still displaying the familiar registration label.
        if duplicates[label] > 1:
            label = '{} ({})'.format(label, field.name)
        columns.append(ExportColumn(
            '{}__{}'.format(prefix, field.name), label,
            (field.name, '{}.{}'.format(prefix, field.name)),
        ))
        if isinstance(field, models.ForeignKey):
            columns.append(ExportColumn(
                '{}__{}_id'.format(prefix, field.name), '{} ID'.format(label),
                ('{}_id'.format(field.name), '{}.{}_id'.format(prefix, field.name)),
            ))
    return columns


def export_columns(fields=None):
    """Return stable, unambiguous headers, retaining identity and photo columns.

    Optional selectors accept field names, ``child__field`` /
    ``registration__field`` names, or the visible header. Unknown legacy SQL
    selectors do not accidentally fall back to exporting every stored field.
    """
    child_columns = _model_columns(Child, 'child', {'id'})
    registration_columns = _model_columns(Registration, 'registration', {'id', 'child'})
    used_labels = {column.label for column in _REQUIRED_COLUMNS + tuple(child_columns)}
    registration_columns = [
        ExportColumn(
            column.key,
            'Registration {}'.format(column.label) if column.label in used_labels else column.label,
            column.aliases,
        ) for column in registration_columns
    ]
    exam_columns = [
        ExportColumn(column.key, 'Exceptional Exam - ' + column.label, column.aliases)
        for column in _model_columns(ExceptionalExam, 'exceptional_exam', {'registration'})
    ]
    optional = [
        ExportColumn('child_full_name', 'Child full name', ('child_fullname',)),
        ExportColumn('date_of_birth', 'Date of birth', ('birthday',)),
    ] + child_columns + registration_columns + exam_columns
    if fields:
        if isinstance(fields, str):
            fields = [fields]
        selected = set(fields)
        optional = [column for column in optional if selected.intersection(
            (column.key, column.label) + column.aliases
        )]
    return list(_REQUIRED_COLUMNS) + optional


def _field_values(instance, model, prefix, excluded):
    values = OrderedDict()
    for field in model._meta.concrete_fields:
        if field.name in excluded:
            continue
        key = '{}__{}'.format(prefix, field.name)
        value = getattr(instance, field.name, None) if instance is not None else None
        if isinstance(field, models.ForeignKey):
            values[key] = _text(value)
            values['{}__{}_id'.format(prefix, field.name)] = _text(
                getattr(instance, field.attname, None) if instance is not None else None
            )
        elif isinstance(field, models.FileField):
            # Storage names are permanent references; signed storage URLs are
            # neither necessary nor appropriate to retain in a document.
            values[key] = _text(value.name if value else None)
        elif field.choices and instance is not None:
            values[key] = _text(getattr(instance, 'get_{}_display'.format(field.name))())
        elif hasattr(field, 'base_field') and field.base_field.choices and value is not None:
            labels = dict(field.base_field.flatchoices)
            values[key] = '; '.join(_text(labels.get(item, item)) for item in value)
        else:
            values[key] = _text(value)
    return values


def _date_of_birth(child):
    """Use the saved birthday components, including incomplete legacy dates."""
    if child is None:
        return ''
    year = _text(child.birthday_year)
    month = _text(child.birthday_month)
    day = _text(child.birthday_day)
    try:
        return datetime.date(int(year), int(month), int(day)).isoformat()
    except (ValueError, TypeError):
        parts = []
        if year and year != '0':
            parts.append(year)
        if month and month != '0':
            parts.append('month: {}'.format(month))
        if day and day != '0':
            parts.append('day: {}'.format(day))
        return ', '.join(parts)


def _load_photo(child):
    """Read the child's own stored image through local or remote storage."""
    if child is None or not child.photo:
        return None, PHOTO_MISSING
    try:
        with child.photo.open('rb') as source:
            data = source.read()
    except Exception:
        # A storage backend may raise its own network or missing-object errors,
        # rather than OSError. Report this child's image as unavailable without
        # leaking storage URLs or replacing it with another child's picture.
        logger.warning('Photograph unavailable for child ID %s', child.pk)
        return None, PHOTO_UNAVAILABLE
    try:
        with Image.open(io.BytesIO(data)) as image:
            # Apply camera orientation before removing metadata and using
            # the same PNG in both types of document.
            image = ImageOps.exif_transpose(image).convert('RGB')
            image.thumbnail((384, 480), Image.Resampling.LANCZOS)
            destination = io.BytesIO()
            image.save(destination, format='PNG')
            return destination.getvalue(), PHOTO_INCLUDED
    except (OSError, ValueError, SyntaxError, Image.DecompressionBombError):
        logger.warning('Photograph unavailable for child ID %s', child.pk)
        return None, PHOTO_UNAVAILABLE


def registration_snapshot(registration):
    """Read one registration's values and its linked child's photograph once."""
    child = registration.child
    photo_data, photo_status = _load_photo(child)
    values = OrderedDict((
        ('registration_id', _text(registration.pk)),
        ('child_id', _text(registration.child_id)),
        ('photograph', _text(child.photo.name if child and child.photo else None)),
        ('photograph_status', photo_status),
        ('child_full_name', ' '.join(
            _text(getattr(child, name, None)).strip()
            for name in ('first_name', 'father_name', 'last_name')
            if child is not None and getattr(child, name, None)
        ) if child is not None else 'Child record unavailable'),
        ('date_of_birth', _date_of_birth(child)),
    ))
    values.update(_field_values(child, Child, 'child', {'id'}))
    values.update(_field_values(registration, Registration, 'registration', {'id', 'child'}))
    exam = None if registration._state.adding else getattr(registration, 'exceptional_exam', None)
    values.update(_field_values(exam, ExceptionalExam, 'exceptional_exam', {'registration'}))
    return RegistrationSnapshot(values, photo_data, photo_status)


def build_children_workbook(registrations, fields=None):
    """Return an XLSX with each photograph contained in its registration row."""
    columns = export_columns(fields)
    destination = io.BytesIO()
    workbook = xlsxwriter.Workbook(destination, {
        'in_memory': True,
        'strings_to_formulas': False,
        'strings_to_urls': False,
    })
    sheet = workbook.add_worksheet('Registered children')
    heading = workbook.add_format({
        'bold': True, 'bg_color': '#164A68', 'font_color': '#FFFFFF',
        'text_wrap': True, 'valign': 'vcenter', 'border': 1,
    })
    body = workbook.add_format({'text_wrap': True, 'valign': 'top'})
    sheet.freeze_panes(1, 2)
    sheet.set_row(0, 42)
    sheet.set_column(0, len(columns) - 1, 24)
    sheet.set_column(0, 1, 15)
    sheet.set_column(2, 2, 20)
    for index, column in enumerate(columns):
        sheet.write_string(0, index, column.label, heading)

    last_row = 0
    for row, registration in enumerate(registrations, start=1):
        if row >= 1048576:
            raise ValueError('The export exceeds the Excel worksheet row limit.')
        last_row = row
        snapshot = registration_snapshot(registration)
        sheet.set_row(row, 108)
        for index, column in enumerate(columns):
            value = snapshot.values.get(column.key, '')
            if column.key == 'photograph':
                value = '' if snapshot.photo_data else snapshot.photo_status
            if len(value) > 32767:
                raise ValueError('A registered value exceeds the Excel cell text limit.')
            # write_string keeps names beginning with '=' or URLs as literal
            # registered values instead of executing formulas or hyperlinks.
            sheet.write_string(row, index, value, body)
        if snapshot.photo_data:
            image_data = io.BytesIO(snapshot.photo_data)
            with Image.open(image_data) as image:
                scale = min(112 / image.width, 136 / image.height, 1)
            sheet.insert_image(row, 2, 'child_{}_registration_{}.png'.format(
                registration.child_id, registration.pk,
            ), {
                'image_data': image_data,
                'x_offset': 4, 'y_offset': 4,
                'x_scale': scale, 'y_scale': scale,
                # Move and size the image with its own row when users sort or
                # filter the workbook. Its full anchor stays inside the row.
                'positioning': 1,
            })
    sheet.autofilter(0, 0, last_row, len(columns) - 1)
    workbook.close()
    return destination.getvalue()


def build_children_csv(registrations, fields=None):
    """Return the same registered values as UTF-8 CSV, with photo references."""
    columns = export_columns(fields)
    destination = io.StringIO(newline='')
    destination.write('\ufeff')
    writer = csv.writer(destination)
    writer.writerow([column.label for column in columns])
    for registration in registrations:
        snapshot = registration_snapshot(registration)
        # CSV has no literal cell type. Escape spreadsheet expression prefixes
        # so opening the file in Excel cannot turn saved text into formulas.
        values = [snapshot.values.get(column.key, '') for column in columns]
        writer.writerow([
            "'" + value if value.lstrip().startswith(('=', '+', '-', '@')) else value
            for value in values
        ])
    return destination.getvalue().encode('utf-8')


def build_examination_card(registration):
    """Return the reference-style landscape card using the saved registration.

    Browser print headers and the empty overflow page in the reference are
    not part of the card. Birthplace is intentionally blank: MSCC children
    have no stored place-of-birth field, and residence is not birthplace.
    """
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    snapshot = registration_snapshot(registration)
    values = snapshot.values
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(210)
    section.page_height = Mm(148)
    section.top_margin = section.bottom_margin = Mm(10)
    section.left_margin = section.right_margin = Mm(12)
    normal = document.styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(16)
    normal.font.bold = True
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = Pt(21)

    borders = OxmlElement('w:pgBorders')
    borders.set(qn('w:offsetFrom'), 'page')
    for side in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement('w:' + side)
        for attribute, value in (
            ('val', 'single'), ('sz', '6'), ('space', '28'), ('color', '000000'),
        ):
            border.set(qn('w:' + attribute), value)
        borders.append(border)
    section._sectPr.insert_element_before(
        borders, 'w:lnNumType', 'w:pgNumType', 'w:cols', 'w:formProt',
        'w:vAlign', 'w:noEndnote', 'w:titlePg', 'w:textDirection', 'w:bidi',
        'w:rtlGutter', 'w:docGrid', 'w:printerSettings', 'w:sectPrChange',
    )

    # The photo floats at the upper left, so the centred details retain the
    # same alignment regardless of the photograph's aspect ratio.
    photo = document.add_paragraph()
    photo.paragraph_format.line_spacing = Pt(1)
    photo.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    photo.paragraph_format.space_after = Pt(0)
    if snapshot.photo_data:
        with Image.open(io.BytesIO(snapshot.photo_data)) as image:
            scale = min(48 / image.width, 42 / image.height)
            width, height = image.width * scale, image.height * scale
        picture = photo.add_run().add_picture(
            io.BytesIO(snapshot.photo_data), width=Mm(width), height=Mm(height),
        )
        anchor = picture._inline
        anchor.tag = qn('wp:anchor')
        for attribute, value in (
            ('simplePos', '0'), ('relativeHeight', '0'), ('behindDoc', '0'),
            ('locked', '0'), ('layoutInCell', '1'), ('allowOverlap', '1'),
        ):
            anchor.set(attribute, value)
        simple_position = OxmlElement('wp:simplePos')
        simple_position.set('x', '0')
        simple_position.set('y', '0')
        anchor.insert(0, simple_position)
        for index, axis in enumerate(('H', 'V'), start=1):
            position = OxmlElement('wp:position' + axis)
            position.set('relativeFrom', 'page')
            offset = OxmlElement('wp:posOffset')
            offset.text = str(Mm(16.5))
            position.append(offset)
            anchor.insert(index, position)
        anchor.insert(4, OxmlElement('wp:wrapNone'))
        anchor.find(qn('wp:docPr')).set(
            'descr', 'Photograph for child {}, registration {}'.format(
                values['child_id'], values['registration_id'],
            ),
        )
    else:
        # Keep missing-image status visible without adding a made-up picture.
        photo.paragraph_format.line_spacing = Pt(12)
        photo.paragraph_format.left_indent = Mm(4.5)
        run = photo.add_run(snapshot.photo_status)
        run.font.size = Pt(9)

    def display_name(instance):
        if instance is None:
            return ''
        return _text(getattr(instance, 'name_en', None) or instance.name).strip()

    def centred_line(text, size=16, arabic=False):
        text = ' '.join(text.split())
        # Longer registered names and labels need smaller type rather than
        # losing characters or spilling the card onto a blank first page.
        if len(text) > 100:
            size = min(size, 9)
        elif len(text) > 68:
            size = min(size, 12)
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        if len(text) > 68:
            paragraph.paragraph_format.line_spacing = Pt(size + 3)
            paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        # Word needs the complex-script font/size and RTL flags explicitly;
        # font.size alone applies only to Latin text in several renderers.
        properties = run._element.get_or_add_rPr()
        fonts = properties.get_or_add_rFonts()
        fonts.set(qn('w:cs'), 'Times New Roman')
        complex_size = OxmlElement('w:szCs')
        complex_size.set(qn('w:val'), str(size * 2))
        properties.insert_element_before(complex_size, 'w:highlight', 'w:u', 'w:rtl')
        run.font.cs_bold = True
        if arabic:
            paragraph._element.get_or_add_pPr().insert_element_before(
                OxmlElement('w:bidi'), 'w:adjustRightInd', 'w:snapToGrid',
                'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
                'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
                'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
                'w:rPr', 'w:sectPr', 'w:pPrChange',
            )
            run.font.rtl = True
        return paragraph

    round_ = registration.round
    round_name = _text(round_.name).strip() if round_ else ''
    if round_ and round_.year and str(round_.year) not in round_name:
        round_name += '/{}'.format(round_.year)
    round_label = round_name if round_name.lower().startswith('round') else 'Round ' + round_name
    title = centred_line(round_label.strip() or 'Round', size=18)
    title.paragraph_format.space_before = Mm(35 if snapshot.photo_data else 31)
    title.paragraph_format.line_spacing = Pt(22)
    title.paragraph_format.space_after = Pt(2)
    heading = centred_line('الامتحان الاستثنائي لطلاب التعليم الغير نظامي', size=11, arabic=True)
    heading.paragraph_format.line_spacing = Pt(15)
    heading.paragraph_format.space_after = Pt(3)

    partner = registration.partner
    ngo = _text((partner.short_name or partner.name) if partner else '').strip()
    child = registration.child
    nationality = display_name(child.nationality) if child else ''
    disability = display_name(child.disability) if child else ''
    for name, extra_key in (('nationality', 'child__nationality_other'),
                            ('disability', 'child__disability_other')):
        label = nationality if name == 'nationality' else disability
        if label.casefold() in ('other', 'others', 'أخرى', 'اخرى', 'آخر', 'اخر') and values[extra_key]:
            label += ': ' + values[extra_key]
        if name == 'nationality':
            nationality = label
        else:
            disability = label
    if disability.casefold() in ('none', 'no', 'no disability', 'لا', 'لا يوجد', 'بدون إعاقة'):
        disability = 'No'
    governorate = display_name(child.governorate) if child else ''
    governorate = governorate or values.get('child__governorate_legacy', '')
    birthday = values['date_of_birth']
    try:
        date = datetime.date.fromisoformat(birthday)
        birthday = '{}/{}/{:02d}'.format(date.day, date.month, date.year % 100)
    except ValueError:
        pass  # Preserve incomplete saved dates instead of fabricating one.

    centred_line('ID: ' + values['registration_id'])
    centred_line('NGO: ' + ngo)
    centred_line('الاسم الثلاثي: ' + values['child_full_name'], arabic=True)
    centred_line('Date of Birth: ' + birthday)
    centred_line('Place of Birth:')
    centred_line('Nationality: ' + nationality)
    centred_line('Governorate: ' + governorate)
    last_line = centred_line('Physical difficulties: ' + disability)
    last_line.paragraph_format.keep_with_next = False
    document.core_properties.title = 'Examination card - Child {}'.format(values['child_id'])
    document.core_properties.subject = 'MSCC / Makani registration {}'.format(values['registration_id'])
    destination = io.BytesIO()
    document.save(destination)
    return destination.getvalue()
